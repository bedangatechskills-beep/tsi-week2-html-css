"""md2docx.py — TSI instructor-guide Markdown -> .docx (session-designer docx pattern, print palette).

Usage:  python md2docx.py <in.md> <out.docx> [--no-cover]

Markdown subset:
  front matter (--- program/label/title/subtitle/duration/footer ---) -> cover page
  '# ' H1 (each starts a new page) · '## ' H2 · '### ' H3 · paragraphs
  '- ' bullets (2 spaces per nesting level) · '1. ' numbered (restarts per list)
  ``` fenced code -> shaded Consolas block · pipe tables (first row = header, navy)
  '> **LABEL:** text' blockquote -> one-cell shaded callout box
  inline **bold**, *italic*, `code`
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

# print palette from tsi-session-designer references/docx-pattern.md
NAVY, BLUE, GREY = '1F3864', '2E74B5', '595959'
CODE_BG, ACCENT_BG, TIP_BG = 'F2F2F2', 'DEEAF6', 'FFF2CC'
ACCENT_LABELS = ('CHECKPOINT', 'ACTIVITY', 'REVEAL', 'SAY', 'DONE')
BODY_FONT, CODE_FONT = 'Calibri', 'Consolas'


def rgb(h):
    return RGBColor.from_string(h)


def shade(el, fill):
    pr = el.get_or_add_tcPr() if el.tag.endswith('}tc') else el.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pr.append(shd)


def para_border(p, sides, color='BFBFBF', sz='4'):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    for side in sides:
        e = OxmlElement(f'w:{side}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), sz)
        e.set(qn('w:space'), '4')
        e.set(qn('w:color'), color)
        bdr.append(e)
    pPr.append(bdr)


def cell_borders(table, color='BFBFBF'):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{side}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), color)
        b.append(e)
    tblPr.append(b)


INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))')


def add_inline(p, text, size=None, color=None, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            add_inline(p, part[2:-2], size=size, color=color, bold=True)
            continue
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            r = p.add_run(part[1:-1]); r.font.name = CODE_FONT
            r._element.rPr.rFonts.set(qn('w:eastAsia'), CODE_FONT)
            r.font.color.rgb = rgb('7F1F2E') if not color else rgb(color)
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            r = p.add_run(part)
        if bold:
            r.bold = True
        if size:
            r.font.size = Pt(size)
        if color and not (part.startswith('`')):
            r.font.color.rgb = rgb(color)
    return p


class Builder:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.left_margin = sec.right_margin = Inches(0.85)
        sec.top_margin = sec.bottom_margin = Inches(0.8)
        st = self.doc.styles['Normal']
        st.font.name = BODY_FONT
        st.element.rPr.rFonts.set(qn('w:eastAsia'), BODY_FONT)
        st.font.size = Pt(10.5)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.line_spacing = 1.1
        for name, size, color in (('Heading 1', 18, NAVY), ('Heading 2', 13.5, BLUE), ('Heading 3', 11.5, GREY)):
            h = self.doc.styles[name]
            h.font.name = BODY_FONT
            h.element.rPr.rFonts.set(qn('w:ascii'), BODY_FONT)
            h.element.rPr.rFonts.set(qn('w:hAnsi'), BODY_FONT)
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = rgb(color)
            h.paragraph_format.space_before = Pt(12 if name != 'Heading 1' else 0)
            h.paragraph_format.space_after = Pt(5)
            h.paragraph_format.keep_with_next = True
        self.first_h1 = True
        self.footer_text = ''

    # ------------------------------------------------------------ page furniture
    def footer(self, text):
        self.footer_text = text
        p = self.doc.sections[0].footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text + '   ·   page ')
        r.font.size = Pt(8.5); r.font.color.rgb = rgb(GREY)
        for kind, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
            run = p.add_run(); run.font.size = Pt(8.5); run.font.color.rgb = rgb(GREY)
            if kind:
                f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), kind); run._r.append(f)
            else:
                t = OxmlElement('w:instrText'); t.set(qn('xml:space'), 'preserve'); t.text = txt; run._r.append(t)

    def cover(self, fm):
        d = self.doc
        for _ in range(5):
            d.add_paragraph()
        def line(text, size, color, bold=False, space=8, italic=False):
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(space)
            r = p.add_run(text); r.font.size = Pt(size); r.font.color.rgb = rgb(color); r.bold = bold; r.italic = italic
        line(fm.get('program', ''), 13, GREY)
        line(fm.get('label', 'Instructor Detail Guide').upper(), 12, BLUE, bold=True, space=26)
        line(fm.get('title', ''), 26, NAVY, bold=True, space=10)
        if fm.get('subtitle'):
            line(fm['subtitle'], 16, BLUE, italic=True, space=24)
        if fm.get('duration'):
            line(fm['duration'], 12, GREY, space=60)
        if fm.get('footer'):
            line(fm['footer'], 11, GREY)
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------ blocks
    def heading(self, level, text):
        h = self.doc.add_heading(level=level)
        if level == 1:
            # page break as a paragraph property: never leaves an empty spill-over page
            h.paragraph_format.page_break_before = not self.first_h1
            self.first_h1 = False
        add_inline(h, text)
        return h

    def paragraph(self, text):
        add_inline(self.doc.add_paragraph(), text)

    def bullet(self, text, level):
        style = 'List Bullet' if level == 0 else 'List Bullet 2' if level == 1 else 'List Bullet 3'
        p = self.doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(2)
        add_inline(p, text)

    def numbered(self, n, text, level):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Inches(0.3 + 0.3 * level); pf.first_line_indent = Inches(-0.25); pf.space_after = Pt(2)
        r = p.add_run(f'{n}.\t'); r.bold = True; r.font.color.rgb = rgb(NAVY)
        pf.tab_stops.add_tab_stop(Inches(0.3 + 0.3 * level))
        add_inline(p, text)

    def code(self, lines):
        lines = lines or ['']
        for i, ln in enumerate(lines):
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
            pf.left_indent = Inches(0.1); pf.keep_with_next = i < len(lines) - 1 and len(lines) <= 25
            shade(p._p, CODE_BG)
            sides = ['left', 'right'] + (['top'] if i == 0 else []) + (['bottom'] if i == len(lines) - 1 else [])
            para_border(p, sides)
            r = p.add_run(ln if ln else ' ')
            r.font.name = CODE_FONT; r._element.rPr.rFonts.set(qn('w:eastAsia'), CODE_FONT)
            r.font.size = Pt(9)
        gap = self.doc.add_paragraph(); gap.paragraph_format.space_after = Pt(4)

    def callout(self, label, text_lines):
        fill = ACCENT_BG if label.upper().rstrip(':').strip() in ACCENT_LABELS else TIP_BG
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.cell(0, 0)
        shade(cell._tc, fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(label.rstrip(':') + ': '); r.bold = True; r.font.color.rgb = rgb(NAVY)
        first = True
        for ln in text_lines:
            if not first:
                if re.match(r'^\s*[-*] ', ln):
                    p = cell.add_paragraph(); p.paragraph_format.left_indent = Inches(0.15); p.paragraph_format.space_after = Pt(1)
                    add_inline(p, '• ' + re.sub(r'^\s*[-*] ', '', ln))
                    continue
                p = cell.add_paragraph(); p.paragraph_format.space_after = Pt(2)
            add_inline(p, ln)
            first = False
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def table(self, rows):
        head, body = rows[0], rows[1:]
        n = len(head)
        t = self.doc.add_table(rows=1 + len(body), cols=n)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        cell_borders(t)
        # column widths in proportion to content length (sqrt-damped), page body = 6.8 in
        import math
        plain = lambda s: re.sub(r'[*`]', '', s)
        lens = [max(len(plain(r[j])) if j < len(r) else 0 for r in rows) for j in range(n)]
        weights = [max(1.2, math.sqrt(max(1, L)) * (1 if L > 12 else 0.85)) for L in lens]
        tot = sum(weights)
        widths = [Inches(6.8 * w / tot) for w in weights]
        for row in t.rows:
            for j, c in enumerate(row.cells):
                c.width = widths[j]
        for j, txt in enumerate(head):
            c = t.cell(0, j); shade(c._tc, NAVY)
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            add_inline(p, txt, size=9.5, color='FFFFFF', bold=True)
            p.paragraph_format.keep_with_next = True   # header never alone at a page bottom
        # repeat header row on each page
        trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)
        for i, row in enumerate(body, start=1):
            for j in range(n):
                txt = row[j] if j < len(row) else ''
                c = t.cell(i, j)
                if i % 2 == 0:
                    shade(c._tc, 'F7F9FC')
                parts = txt.split('<br>')
                p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
                add_inline(p, parts[0], size=9.5)
                for extra in parts[1:]:
                    p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0)
                    add_inline(p, extra, size=9.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)


def split_row(line):
    line = line.strip()
    if line.startswith('|'):
        line = line[1:]
    if line.endswith('|') and not line.endswith('\\|'):
        line = line[:-1]
    cells, cur, in_code = [], '', False
    i = 0
    while i < len(line):
        ch = line[i]
        if ch == '\\' and i + 1 < len(line) and line[i + 1] == '|':
            cur += '|'; i += 2; continue
        if ch == '`':
            in_code = not in_code
        if ch == '|' and not in_code:
            cells.append(cur.strip()); cur = ''
        else:
            cur += ch
        i += 1
    cells.append(cur.strip())
    return cells


def convert(md_path, out_path, cover=True):
    with open(md_path, encoding='utf-8') as f:
        lines = f.read().replace('\r\n', '\n').split('\n')
    fm = {}
    if lines and lines[0].strip() == '---':
        end = lines.index('---', 1)
        for ln in lines[1:end]:
            if ':' in ln:
                k, v = ln.split(':', 1); fm[k.strip()] = v.strip()
        lines = lines[end + 1:]
    B = Builder()
    if fm.get('footer'):
        B.footer(fm.get('title', '') or fm['footer'])
    if cover and fm:
        B.cover(fm)
    i, num = 0, {}
    para_buf = []

    def flush():
        if para_buf:
            B.paragraph(' '.join(s.strip() for s in para_buf)); para_buf.clear()

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if s.startswith('```'):
            flush(); i += 1; buf = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                buf.append(lines[i].rstrip('\n')); i += 1
            B.code(buf); i += 1; continue
        m = re.match(r'^(#{1,3}) (.*)', ln)
        if m:
            flush(); B.heading(len(m.group(1)), m.group(2).strip()); i += 1; num.clear(); continue
        if s.startswith('|'):
            flush(); rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                r = split_row(lines[i])
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in r if c):
                    rows.append(r)
                i += 1
            B.table(rows); continue
        if s.startswith('>'):
            flush(); buf = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                buf.append(re.sub(r'^\s*> ?', '', lines[i])); i += 1
            m2 = re.match(r'^\*\*(.+?)\*\*\s*(.*)', buf[0])
            label, first = (m2.group(1), m2.group(2)) if m2 else ('NOTE', buf[0])
            B.callout(label, [first] + buf[1:]); continue
        m = re.match(r'^(\s*)[-*] (.*)', ln)
        if m:
            flush(); B.bullet(m.group(2), len(m.group(1)) // 2); i += 1; continue
        m = re.match(r'^(\s*)(\d+)\. (.*)', ln)
        if m:
            flush(); lvl = len(m.group(1)) // 2
            num[lvl] = num.get(lvl, 0) + 1
            for k in [k for k in num if k > lvl]:
                del num[k]
            B.numbered(num[lvl], m.group(3), lvl); i += 1; continue
        if not s:
            flush(); i += 1
            continue
        if s == '<!-- pagebreak -->':
            flush(); B.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE); i += 1; continue
        para_buf.append(ln); num.clear(); i += 1
    flush()
    B.doc.save(out_path)
    print(f'{out_path}: written')


if __name__ == '__main__':
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    convert(args[0], args[1], cover='--no-cover' not in sys.argv)
