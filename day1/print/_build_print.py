"""Build the print copies (docx in the TSI docx pattern, then PDF) for one day.

Usage:
    python _build_print.py --day 1|2 [--root <week2>] [--only notes,homework,booklet,keys] [--no-pdf]

Outputs (in week2/dayN/print/):
    W2DN-Student-Notes.docx/.pdf        notes/README.md + notes 01-11
    W2DN-Homework-Sheet.docx/.pdf       homework sheet + HW1/HW2/HW4 READMEs (+ D2-HW2 CHECKLIST) + HW3
    W2DN-Classwork-Booklet.docx/.pdf    every classwork README in id order (D?-00 ... then FF1-FF6), no solutions
    instructor/W2DN-Answer-Keys.docx/.pdf   every BUGS.md (+ answer files) + the HW3 answer key, instructor only

Markdown subset: # headings, paragraphs, **bold**, *italic*, `code`, [text](url), - / 1. lists (nested by indent),
- [ ] task boxes, | tables |, ``` code fences (Courier New, shaded), > quotes (callout box), --- rules.
PDF: Word (COM via PowerShell); fallback LibreOffice (soffice --headless).
"""
import argparse, glob, os, re, shutil, subprocess, sys, tempfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY, BLUE, GREY = "1F3864", "2E74B5", "595959"
CODE_BG, ACCENT_BG, TIP_BG = "F2F2F2", "DEEAF6", "FFF2CC"
BODY_FONT, CODE_FONT = "Calibri", "Courier New"
PROGRAM = "AI-Powered Full Stack Developer Program"
INSTITUTE = "Tech Skills Institute (TSI) Nepal · Kathmandu"
DAY_TITLE = {1: "How the Web Works + HTML Core", 2: "Semantic HTML, Accessibility, Media → CSS Basics"}


# ----------------------------------------------------------------------------- low-level helpers
def rgb(h):
    return RGBColor.from_string(h)


def shade(el_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el_pr.append(shd)


def para_border(p, sides=("top", "bottom", "left", "right"), color="BFBFBF", sz="4", space="4"):
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    for s in sides:
        e = OxmlElement(f"w:{s}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), sz)
        e.set(qn("w:space"), space)
        e.set(qn("w:color"), color)
        bdr.append(e)
    ppr.append(bdr)


def cell_shade(cell, fill):
    shade(cell._tc.get_or_add_tcPr(), fill)


def set_cell_margins(table, top=60, bottom=60, left=90, right=90):
    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblpr.append(mar)


def keep_with_next(p):
    p.paragraph_format.keep_with_next = True


def add_field(run, instr):
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1); run._r.append(it); run._r.append(f2)


INLINE = re.compile(r"(\*\*(?:`[^`]*`|[^*`])+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|(?<![\w*])\*(?:`[^`]*`|[^*\s`])(?:`[^`]*`|[^*`])*?\*(?![\w*]))")


def add_inline(p, text, size=None, color=None, bold=False, italic=False):
    """Add markdown inline text to paragraph p."""
    for part in INLINE.split(text):
        if not part:
            continue
        b, i, mono, t = bold, italic, False, part
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            add_inline(p, part[2:-2], size, color, True, italic)
            continue
        elif part.startswith("`") and part.endswith("`"):
            mono, t = True, part[1:-1]
        elif part.startswith("[") and "](" in part:
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            t = m.group(1)
            if m.group(2).startswith("http") and m.group(2) != t:
                t = f"{t} ({m.group(2)})"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            add_inline(p, part[1:-1], size, color, bold, True)
            continue
        # nested bold inside italics etc. are rare; strip stray markers
        if not mono:
            t = t.replace("\\|", "|").replace("<br>", "\n")  # <br> in a table cell = line break; kept in `code`
        r = p.add_run(t)
        r.bold, r.italic = b, i
        if mono:
            r.font.name = CODE_FONT
            r._element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
            r.font.size = Pt((size or 10.5) - 1)
            r.font.color.rgb = rgb("7A0010")
        else:
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = rgb(color)
    return p


# ----------------------------------------------------------------------------- document builder
class Builder:
    def __init__(self, label, title, day, instructor=False):
        self.doc = Document()
        self.lists = 0
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
        st = self.doc.styles["Normal"]
        st.font.name, st.font.size = BODY_FONT, Pt(10.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.12
        for lvl, (sz, col) in {1: (18, NAVY), 2: (14, NAVY), 3: (12, BLUE), 4: (11, GREY)}.items():
            hs = self.doc.styles[f"Heading {lvl}"]
            hs.font.name, hs.font.size, hs.font.bold = BODY_FONT, Pt(sz), True
            hs.font.color.rgb = rgb(col)
            hs.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
            rfonts = hs.element.rPr.rFonts
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rfonts.get(qn(a)) is not None:
                    del rfonts.attrib[qn(a)]
            hs.paragraph_format.space_before = Pt({1: 6, 2: 14, 3: 10, 4: 8}[lvl])
            hs.paragraph_format.space_after = Pt(4)
        self.day, self.label, self.title = day, label, title
        self._header_footer(instructor)
        self._cover(instructor)

    def _header_footer(self, instructor):
        sec = self.doc.sections[0]
        sec.different_first_page_header_footer = True
        hp = sec.header.paragraphs[0]
        hp.text = ""
        r = hp.add_run(f"TSI · Week 2 Day {self.day} · {self.label}" + ("  ·  INSTRUCTOR ONLY" if instructor else ""))
        r.font.size, r.font.color.rgb = Pt(8.5), rgb(GREY)
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("You may only ship code you can explain, line by line.   ·   page ")
        r.font.size, r.font.color.rgb = Pt(8.5), rgb(GREY)
        r2 = fp.add_run()
        r2.font.size = Pt(8.5)
        add_field(r2, "PAGE")

    def _cover(self, instructor):
        d = self.doc
        for _ in range(6):
            d.add_paragraph()
        def c(text, size, color, bold=False, after=6):
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(after)
            r = p.add_run(text)
            r.font.size, r.font.color.rgb, r.bold = Pt(size), rgb(color), bold
            return p
        c(PROGRAM, 13, BLUE, True)
        c(self.label.upper() + ("  ·  INSTRUCTOR ONLY" if instructor else ""), 11, GREY, True, 30)
        c(f"Week 2 · Day {self.day}", 16, NAVY, True)
        c(DAY_TITLE[self.day], 24, NAVY, True, 24)
        c(self.title, 12, GREY, False, 60)
        c("Momo House · TSI Canteen · Aarav · Sunita · Didi" + (" · Rajan" if self.day == 2 else ""), 10.5, GREY)
        c("Stage 1, AI as Tutor: AI may explain; you type every line.", 10.5, GREY, False, 80)
        c(INSTITUTE, 10.5, NAVY, True)
        self.page_break()

    # -- blocks
    def page_break(self):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)

    def heading(self, text, level):
        level = min(level, 4)
        h = self.doc.add_heading(level=level)
        add_inline(h, text)
        keep_with_next(h)
        return h

    def paragraph(self, text):
        p = self.doc.add_paragraph()
        add_inline(p, text)
        if text.rstrip().endswith(":") and len(text) < 200:
            keep_with_next(p)
        return p

    def code(self, lines, broken=False):
        if not lines:
            lines = [""]
        n = len(lines)
        longest = max(len(ln.replace("	", "  ")) for ln in lines)
        size = 8.5 if longest <= 88 else max(6.5, round(8.5 * 88 / longest * 2) / 2)  # 17 cm fits ~88 chars at 8.5 pt
        for k, ln in enumerate(lines):
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = Pt(4 if k == 0 else 0)
            pf.space_after = Pt(6 if k == n - 1 else 0)
            pf.line_spacing = 1.0
            pf.left_indent = Cm(0.15)
            pf.right_indent = Cm(0.15)
            if n <= 25 and k < n - 1:
                pf.keep_with_next = True
            shade(p._p.get_or_add_pPr(), "FBEAEA" if broken else CODE_BG)
            sides = ["left", "right"] + (["top"] if k == 0 else []) + (["bottom"] if k == n - 1 else [])
            para_border(p, sides, color="C00000" if broken else "BFBFBF")
            r = p.add_run(ln.replace("\t", "  ") if ln else " ")
            r.font.name = CODE_FONT
            r._element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
            r.font.size = Pt(size)
        if broken:
            p = self.doc.paragraphs[-n]
            # label above the block
            lab = p.insert_paragraph_before()
            rr = lab.add_run("✗ Broken on purpose: find what is wrong")
            rr.font.size, rr.italic, rr.font.color.rgb = Pt(8.5), True, rgb("C00000")
            lab.paragraph_format.space_after = Pt(0)
            keep_with_next(lab)

    def callout(self, lines):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.rows[0].cells[0]
        cell_shade(cell, TIP_BG)
        set_cell_margins(t, 80, 80, 140, 140)
        first = True
        for ln in lines:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            m = re.match(r"^\s*[-*]\s+(.*)", ln)
            add_inline(p, ("• " + m.group(1)) if m else ln)
            p.paragraph_format.space_after = Pt(2)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def bullet(self, text, level, ordered=False, num=None):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.6 + 0.6 * level)
        pf.first_line_indent = Cm(-0.45)
        pf.space_after = Pt(2)
        task = re.match(r"^\[( |x|X)\]\s+(.*)", text)
        if task:
            marker, text = ("☑ " if task.group(1).lower() == "x" else "☐ "), task.group(2)
        elif ordered:
            marker = f"{num}. "
        else:
            marker = ["• ", "– ", "· "][min(level, 2)]
        r = p.add_run(marker)
        r.font.color.rgb = rgb(BLUE)
        r.bold = ordered
        add_inline(p, text)
        return p

    def table(self, rows):
        ncol = max(len(r) for r in rows)
        rows = [r + [""] * (ncol - len(r)) for r in rows]
        t = self.doc.add_table(rows=len(rows), cols=ncol)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_cell_margins(t, 30, 30, 70, 70)
        fs = 9 if ncol <= 4 else 8.5 if ncol <= 6 else 7.5
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.rows[i].cells[j]
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                if i == 0:
                    cell_shade(cell, NAVY)
                    add_inline(p, val, size=fs, color="FFFFFF", bold=True)
                else:
                    if i % 2 == 0:
                        cell_shade(cell, "F7F9FC")
                    add_inline(p, val, size=fs)
        # column widths from content: long-text columns get more room, short ones keep a readable minimum
        def plain(v):
            return re.sub(r"[`*]|\[([^\]]*)\]\([^)]*\)", lambda m: m.group(1) or "", v)
        weights = []
        for j in range(ncol):
            lens = [len(plain(r[j])) for r in rows]
            longest_word = max((len(w) for r in rows for w in plain(r[j]).split()), default=4)
            weights.append(max(min(sum(lens) / len(lens), 45), longest_word * 1.15, 5))
        total_w = Cm(17.0)
        mins = Cm(1.7)
        raw = [total_w * w / sum(weights) for w in weights]
        longest = [max((len(w) for r in rows for w in plain(r[j]).split()), default=4) for j in range(ncol)]
        floor = [max(mins, Cm(min(0.4 + 0.2 * fs / 9 * lw, 4.0))) for lw in longest]
        raw = [max(x, f) for x, f in zip(raw, floor)]
        for _ in range(5):  # shrink only the columns that are above their floor
            extra = sum(raw) - total_w
            flex = [j for j in range(ncol) if raw[j] > floor[j]]
            if extra <= 0 or not flex:
                break
            room = sum(raw[j] - floor[j] for j in flex)
            for j in flex:
                raw[j] -= min(raw[j] - floor[j], extra * (raw[j] - floor[j]) / room)
        scale = min(1.0, total_w / sum(raw))
        t.autofit = False
        for j in range(ncol):
            wj = int(raw[j] * scale)
            for i in range(len(rows)):
                t.rows[i].cells[j].width = wj
        # repeat header row
        trpr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trpr.append(th)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def rule(self):
        p = self.doc.add_paragraph()
        para_border(p, ("bottom",), color="BFBFBF", sz="6", space="1")
        p.paragraph_format.space_after = Pt(6)

    def source_label(self, text):
        p = self.doc.add_paragraph()
        r = p.add_run(text)
        r.font.size, r.font.color.rgb, r.italic = Pt(8.5), rgb(GREY), True
        p.paragraph_format.space_after = Pt(2)

    # -- markdown
    def markdown(self, md, shift=0):
        lines = md.replace("\r\n", "\n").split("\n")
        i, n = 0, len(lines)
        para_buf = []
        num_stack = {}

        def flush():
            if para_buf:
                self.paragraph(" ".join(s.strip() for s in para_buf))
                para_buf.clear()

        while i < n:
            ln = lines[i]
            s = ln.strip()
            if s.startswith("```"):
                flush()
                info = s[3:].strip().lower()
                j = i + 1
                block = []
                while j < n and not lines[j].strip().startswith("```"):
                    block.append(lines[j])
                    j += 1
                self.code(block, broken="broken" in info)
                i = j + 1
                continue
            if not s:
                flush(); num_stack.clear(); i += 1; continue
            m = re.match(r"^(#{1,6})\s+(.*)", s)
            if m:
                flush(); num_stack.clear()
                self.heading(m.group(2).strip(), len(m.group(1)) + shift)
                i += 1; continue
            if re.match(r"^(-{3,}|\*{3,}|_{3,})$", s):
                flush(); self.rule(); i += 1; continue
            if s.startswith("|"):
                flush()
                rows = []
                while i < n and lines[i].strip().startswith("|"):
                    row = lines[i].strip()
                    cells = [c.strip() for c in re.split(r"(?<!\\)\|", row.strip("|"))]
                    if not all(re.match(r"^:?-{2,}:?$", c) for c in cells if c):
                        rows.append(cells)
                    i += 1
                if rows:
                    self.table(rows)
                continue
            if s.startswith(">"):
                flush()
                q = []
                while i < n and lines[i].strip().startswith(">"):
                    q.append(lines[i].strip()[1:].strip())
                    i += 1
                # join soft-wrapped lines into paragraphs
                merged, cur = [], []
                for x in q:
                    if not x or re.match(r"^[-*]\s", x):
                        if cur: merged.append(" ".join(cur)); cur = []
                        if x: merged.append(x)
                    else:
                        cur.append(x)
                if cur: merged.append(" ".join(cur))
                self.callout(merged)
                continue
            lm = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)", ln)
            if lm:
                flush()
                indent = len(lm.group(1).replace("\t", "    "))
                ordered = lm.group(2)[0].isdigit()
                text = lm.group(3)
                # continuation lines (indented, not a new item)
                k = i + 1
                while k < n and lines[k].strip() and not re.match(r"^\s*([-*+]|\d+[.)])\s+", lines[k]) \
                        and lines[k].startswith(" ") and not lines[k].strip().startswith("```") \
                        and not lines[k].strip().startswith("|"):
                    text += " " + lines[k].strip()
                    k += 1
                lvl = 0 if indent == 0 else 1 if indent <= 4 else 2
                if ordered:
                    num = int(re.match(r"\d+", lm.group(2)).group())
                    self.bullet(text, lvl, True, num)
                else:
                    self.bullet(text, lvl)
                i = k
                continue
            para_buf.append(ln)
            i += 1
        flush()

    def save(self, path):
        os.makedirs(os.path.dirname(path), exist_ok=True)
        self.doc.save(path)


# ----------------------------------------------------------------------------- pdf
def to_pdf(docx_path):
    pdf = os.path.splitext(docx_path)[0] + ".pdf"
    ps = (
        "$ErrorActionPreference='Stop';"
        "$w=New-Object -ComObject Word.Application; $w.Visible=$false; $w.DisplayAlerts=0;"
        f"try {{ $d=$w.Documents.Open('{docx_path}', $false, $true);"
        f"$d.ExportAsFixedFormat('{pdf}', 17); $d.Close(0) }} finally {{ $w.Quit() }}"
    )
    try:
        p = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", ps],
                           capture_output=True, text=True, timeout=300)
        if p.returncode == 0 and os.path.exists(pdf):
            return pdf, "Word"
        err = p.stderr.strip()[:300]
    except Exception as e:  # noqa
        err = str(e)
    soffice = shutil.which("soffice") or r"C:\Program Files\LibreOffice\program\soffice.exe"
    if os.path.exists(soffice):
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", os.path.dirname(docx_path), docx_path],
                       capture_output=True, timeout=300)
        if os.path.exists(pdf):
            return pdf, "LibreOffice"
    raise RuntimeError(f"PDF conversion failed for {docx_path}: {err}")


# ----------------------------------------------------------------------------- sources
def read(path):
    return open(path, encoding="utf-8").read()


def strip_title(md):
    """Return (first H1 text, rest) so the file title becomes a level-1 heading we control."""
    m = re.match(r"\s*#\s+(.+)\n", md)
    if m:
        return m.group(1).strip(), md[m.end():]
    return None, md


def id_key(name):
    m = re.match(r"D\d-(FF|HW)?(\d+)", name)
    if not m:
        return (9, name)
    return ({None: 0, "FF": 1, "HW": 2}[m.group(1)], int(m.group(2)))


def add_file(b, path, root, first=False, label=None):
    if not first:
        b.page_break()
    title, body = strip_title(read(path))
    b.heading(title or label or os.path.basename(path), 1)
    b.source_label("File: " + os.path.relpath(path, root).replace(os.sep, "/"))
    b.markdown(body, shift=0)


def build_notes(day_dir, day, out):
    files = [os.path.join(day_dir, "notes", "README.md")] + sorted(glob.glob(os.path.join(day_dir, "notes", "[0-9][0-9]-*.md")))
    files = [f for f in files if os.path.exists(f)]
    if not files:
        return None, "notes missing"
    b = Builder("Student Notes", f"Theory notes 01–{len(files) - 1:02d}: read them before homework, keep them for interviews", day)
    for k, f in enumerate(files):
        add_file(b, f, day_dir, first=(k == 0))
    b.save(out)
    return out, f"{len(files)} files"


def build_homework(day_dir, day, out):
    hw = os.path.join(day_dir, "homework")
    files = [os.path.join(hw, f"W2D{day}-Homework-Sheet.md")]
    for d in sorted(glob.glob(os.path.join(hw, f"D{day}-HW*")), key=lambda p: id_key(os.path.basename(p))):
        if os.path.isdir(d):
            files += [p for p in (os.path.join(d, "README.md"), os.path.join(d, "CHECKLIST.md")) if os.path.exists(p)]
        elif d.endswith(".md"):
            files.append(d)
    files = [f for f in files if os.path.exists(f)]
    b = Builder("Homework Sheet", "Homework HW1–HW4: tasks, time boxes, submit checklist, self-check and interview prep", day)
    for k, f in enumerate(files):
        add_file(b, f, day_dir, first=(k == 0))
    b.save(out)
    return out, f"{len(files)} files"


def classwork_dirs(day_dir, day):
    cw = os.path.join(day_dir, "classwork")
    main = [d for d in glob.glob(os.path.join(cw, f"D{day}-[0-9][0-9]-*")) if os.path.isdir(d)]
    ff = [d for d in glob.glob(os.path.join(cw, "fast-finishers", f"D{day}-FF*")) if os.path.isdir(d)]
    return sorted(main, key=lambda p: id_key(os.path.basename(p))), sorted(ff, key=lambda p: id_key(os.path.basename(p)))


def build_booklet(day_dir, day, out):
    main, ff = classwork_dirs(day_dir, day)
    files = [os.path.join(d, "README.md") for d in main]
    ffr = os.path.join(day_dir, "classwork", "fast-finishers", "README.md")
    files += ([ffr] if os.path.exists(ffr) else []) + [os.path.join(d, "README.md") for d in ff]
    files = [f for f in files if os.path.exists(f)]
    if not files:
        return None, "classwork READMEs missing"
    b = Builder("Classwork Booklet", f"Every classwork task in order: D{day}-00 … then fast finishers FF1–FF6 (no solutions)", day)
    for k, f in enumerate(files):
        add_file(b, f, day_dir, first=(k == 0))
    b.save(out)
    return out, f"{len(files)} READMEs"


def build_keys(day_dir, day, out):
    main, ff = classwork_dirs(day_dir, day)
    files = []
    for d in main + ff:
        for name in ("BUGS.md", "ANSWERS.md", "README-answer.md", "sample-findings.md"):
            p = os.path.join(d, "solution", name)
            if os.path.exists(p):
                files.append((p, os.path.basename(d)))
    sol = os.path.join(day_dir, "homework", "_solutions")
    for d in sorted(glob.glob(os.path.join(sol, f"D{day}-HW*")), key=lambda p: id_key(os.path.basename(p))):
        if os.path.isdir(d):
            p = os.path.join(d, "BUGS.md")
            if os.path.exists(p):
                files.append((p, os.path.basename(d)))
        elif d.endswith(".md"):
            files.append((d, None))
    if not files:
        return None, "no BUGS.md found"
    b = Builder("Answer Keys", "Every bug-hunt's BUGS.md, the other answer files, and the HW3 answer key", day, instructor=True)
    for k, (f, folder) in enumerate(files):
        if k:
            b.page_break()
        md = read(f)
        title, body = strip_title(md)
        head = f"{folder} · {os.path.basename(f)}" if folder else (title or os.path.basename(f))
        b.heading(head, 1)
        if folder and title:
            b.heading(title, 2)
        b.source_label("File: " + os.path.relpath(f, day_dir).replace(os.sep, "/"))
        b.markdown(body)
    b.save(out)
    return out, f"{len(files)} files"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--day", type=int, required=True, choices=(1, 2))
    ap.add_argument("--root", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
    ap.add_argument("--only", default="notes,homework,booklet,keys")
    ap.add_argument("--no-pdf", action="store_true")
    a = ap.parse_args()
    root = os.path.abspath(a.root)
    if os.path.basename(root).startswith("day"):
        root = os.path.dirname(root)
    day_dir = os.path.join(root, f"day{a.day}")
    pr = os.path.join(day_dir, "print")
    jobs = {
        "notes": (build_notes, os.path.join(pr, f"W2D{a.day}-Student-Notes.docx")),
        "homework": (build_homework, os.path.join(pr, f"W2D{a.day}-Homework-Sheet.docx")),
        "booklet": (build_booklet, os.path.join(pr, f"W2D{a.day}-Classwork-Booklet.docx")),
        "keys": (build_keys, os.path.join(pr, "instructor", f"W2D{a.day}-Answer-Keys.docx")),
    }
    rc = 0
    for name in [x.strip() for x in a.only.split(",") if x.strip()]:
        fn, out = jobs[name]
        path, info = fn(day_dir, a.day, out)
        if not path:
            print(f"[{name}] SKIPPED: {info}")
            rc = 1
            continue
        msg = f"[{name}] {os.path.relpath(path, root)} ({info})"
        if not a.no_pdf:
            pdf, how = to_pdf(os.path.abspath(path))
            msg += f" -> {os.path.basename(pdf)} via {how}"
        print(msg)
    sys.exit(rc)


if __name__ == "__main__":
    main()
